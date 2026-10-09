import io
import json
import sys
import time
from pathlib import Path

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app import Settings, create_app


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(Settings(tmp_path))) as client:
        yield client


def make_case(client, name='Test case'):
    response = client.post('/api/cases', json={'name': name})
    assert response.status_code == 201
    return response.json()['id']


def upload(client, case, name='evidence.txt', content=b'On 12 March 2025 a payment of GBP 500 was agreed by email.'):
    return client.post('/api/cases/' + case + '/documents', files={'file': (name, content)})


def wait_for_tasks(client, case):
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        tasks = client.get('/api/cases/' + case + '/tasks').json()
        if tasks and all(task['status'] in {'completed', 'failed'} for task in tasks):
            return tasks
        time.sleep(.02)
    pytest.fail('Review tasks did not finish.')


def test_upload_parallel_agents_citations_and_export(client):
    case = make_case(client)
    doc = upload(client, case).json()
    agents = client.get('/api/cases/' + case + '/agents').json()
    assert len(agents) == 4
    review = client.post('/api/cases/' + case + '/reviews', json={'agent_ids': [agent['id'] for agent in agents[:3]]})
    assert review.status_code == 202
    tasks = wait_for_tasks(client, case)
    assert len(tasks) == 3
    assert all(task['status'] == 'completed' and task['findings'] for task in tasks)
    for task in tasks:
        for finding in task['findings']:
            assert finding['document_id'] == doc['id']
            evidence = client.get('/api/cases/' + case + '/documents/' + doc['id']).json()
            passage = next(p for p in evidence['passages'] if p['id'] == finding['passage_id'])
            assert passage['text'] == finding['quote']
    exported = client.get('/api/cases/' + case + '/export').json()
    assert len(exported['reviews']) == 3
    assert client.get('/api/cases/' + case + '/documents/' + doc['id'] + '/download').content.startswith(b'On 12 March')


def test_chat_search_and_no_match_are_honest_and_persistent(client):
    case = make_case(client)
    upload(client, case)
    result = client.post('/api/cases/' + case + '/chat', json={'question': 'Find the payment'}).json()
    assert result['mode'] == 'evidence'
    assert result['citations'] and '500' in result['citations'][0]['quote']
    none = client.post('/api/cases/' + case + '/chat', json={'question': 'zebras'}).json()
    assert not none['citations']
    assert 'does not establish' in none['text']
    assert len(client.get('/api/cases/' + case + '/messages').json()) == 4


def test_case_isolation_and_snapshot_selection(client):
    first, second = make_case(client, 'First'), make_case(client, 'Second')
    first_doc = upload(client, first).json()
    second_doc = upload(client, second, content=b'A private settlement payment.').json()
    assert client.get('/api/cases/' + first + '/documents/' + second_doc['id']).status_code == 404
    agents = client.get('/api/cases/' + second + '/agents').json()
    assert client.post('/api/cases/' + first + '/reviews', json={'agent_ids': [agents[0]['id']]}).status_code == 404
    assert client.post('/api/cases/' + first + '/chat', json={'question': 'payment', 'document_ids': [second_doc['id']]}).status_code == 404
    assert client.post('/api/cases/' + first + '/reviews', json={'agent_ids': [client.get('/api/cases/' + first + '/agents').json()[0]['id']], 'document_ids': []}).status_code == 422
    later = upload(client, first, 'later.txt', b'An unrelated payment of 9000.').json()
    response = client.post('/api/cases/' + first + '/chat', json={'question': 'payment', 'document_ids': [first_doc['id']]}).json()
    assert all(c['document_id'] != later['id'] for c in response['citations'])


def test_custom_agent_edit_and_no_findings(client):
    case = make_case(client)
    upload(client, case)
    data = {'name': 'Repairs', 'instruction': 'Look for repairs and leaks.', 'keywords': 'repair,leak'}
    agent = client.post('/api/cases/' + case + '/agents', json=data).json()
    client.post('/api/cases/' + case + '/reviews', json={'agent_ids': [agent['id']]})
    task = wait_for_tasks(client, case)[0]
    assert task['status'] == 'completed' and not task['findings']
    data['keywords'] = 'payment'
    assert client.put('/api/cases/' + case + '/agents/' + agent['id'], json=data).status_code == 200
    assert next(a for a in client.get('/api/cases/' + case + '/agents').json() if a['id'] == agent['id'])['keywords'] == 'payment'


def test_document_formats_and_upload_failures(client):
    case = make_case(client)
    docx = Document(); docx.add_paragraph('An agreement for a payment of 250.')
    file = io.BytesIO(); docx.save(file)
    assert upload(client, case, 'agreement.docx', file.getvalue()).status_code == 201
    book = Workbook(); book.active.append(['Invoice', 'Amount']); book.active.append(['Invoice 1', 250])
    file = io.BytesIO(); book.save(file)
    assert upload(client, case, 'invoices.xlsx', file.getvalue()).status_code == 201
    assert upload(client, case, 'email.eml', b'From: a@example.com\nTo: b@example.com\nSubject: Payment\nContent-Type: text/plain; charset=utf-8\n\nPayment was agreed.').status_code == 201
    assert upload(client, case, 'values.csv', b'Invoice,Amount\nA,250').status_code == 201
    assert upload(client, case, 'scan.pdf', b'Not a PDF').status_code == 422
    assert upload(client, case, 'script.exe', b'anything').status_code == 415
    assert upload(client, case, 'empty.txt', b'').status_code == 422
    assert upload(client, case, 'binary.txt', b'\x00data').status_code == 422
    assert upload(client, case, 'too-big.txt', b'x' * (25 * 1024 * 1024 + 1)).status_code == 413
    assert client.post('/api/cases', json={'name': '   '}).status_code == 422
    assert client.post('/api/cases/' + case + '/chat', json={'question': '  '}).status_code == 422


def test_readable_and_scanned_pdf(client):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    case = make_case(client)
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject(); stream.set_data(b'BT /F1 12 Tf 50 750 Td (Payment of 500 agreed on 12 March 2025.) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    file = io.BytesIO(); writer.write(file)
    response = upload(client, case, 'payment.pdf', file.getvalue())
    assert response.status_code == 201
    doc = client.get('/api/cases/' + case + '/documents/' + response.json()['id']).json()
    assert doc['passages'][0]['page'] == 1 and 'Payment of 500' in doc['passages'][0]['text']
    blank = PdfWriter(); blank.add_blank_page(width=612, height=792)
    file = io.BytesIO(); blank.write(file)
    response = upload(client, case, 'scan.pdf', file.getvalue())
    assert response.status_code == 422 and 'OCR' in response.json()['detail']


def test_external_ai_requires_per_action_consent_and_validates_citations(tmp_path, monkeypatch):
    settings = Settings(tmp_path, 'https://example.test/v1', 'test-model', 'test-placeholder')
    real_client = httpx.Client
    sent = []

    def respond(request):
        sent.append(json.loads(request.content))
        prompt = sent[-1]['messages'][1]['content']
        evidence = json.loads(prompt.split('Evidence:\n')[1])
        content = {'answer': 'The payment is documented.', 'passage_ids': [evidence[0]['passage_id']]}
        return httpx.Response(200, json={'choices': [{'message': {'content': json.dumps(content)}}]})

    monkeypatch.setattr('app.httpx.Client', lambda **kwargs: real_client(transport=httpx.MockTransport(respond), **kwargs))
    with TestClient(create_app(settings)) as client:
        case = make_case(client); doc = upload(client, case).json()
        denied = client.post('/api/cases/' + case + '/chat', json={'question': 'payment', 'use_ai': True})
        assert denied.status_code == 409 and not sent
        allowed = client.post('/api/cases/' + case + '/chat', json={'question': 'payment', 'use_ai': True, 'approved_external': True, 'document_ids': [doc['id']]})
        assert allowed.status_code == 200
        assert allowed.json()['citations'][0]['document_id'] == doc['id']
        assert len(sent) == 1
        assert client.post('/api/cases/' + case + '/chat', json={'question': 'payment', 'use_ai': True}).status_code == 409
        assert len(sent) == 1
        with client.app.state.store.db() as db:
            approval = db.execute('SELECT * FROM approvals').fetchone()
            assert json.loads(approval['document_ids']) == [doc['id']]
        assert 'test-placeholder' not in json.dumps(client.get('/api/settings').json())


def test_provider_faults_and_fabricated_citations_do_not_save_answers(tmp_path, monkeypatch):
    settings = Settings(tmp_path, 'https://example.test/v1', 'test-model', 'test-placeholder')
    real_client = httpx.Client
    responses = iter([
        httpx.Response(401, json={'error': 'private provider message'}),
        httpx.Response(200, json={'choices': [{'message': {'content': '{"answer":"Unsupported","passage_ids":["invented"]}'}}]}),
    ])
    monkeypatch.setattr('app.httpx.Client', lambda **kwargs: real_client(transport=httpx.MockTransport(lambda req: next(responses)), **kwargs))
    with TestClient(create_app(settings)) as client:
        case = make_case(client); upload(client, case)
        for expected in ['HTTP 401', 'unsupported citation']:
            response = client.post('/api/cases/' + case + '/chat', json={'question': 'payment', 'use_ai': True, 'approved_external': True})
            assert response.status_code == 502 and expected in response.json()['detail']
        assert client.get('/api/cases/' + case + '/messages').json() == []


def test_ai_agent_results_and_failures(tmp_path, monkeypatch):
    settings = Settings(tmp_path, 'https://example.test/v1', 'test-model', 'test-placeholder')
    real_client = httpx.Client
    invalid = False

    def respond(request):
        evidence = json.loads(json.loads(request.content)['messages'][1]['content'].split('Evidence:\n')[1])
        finding = {'passage_id': 'invented' if invalid else evidence[0]['passage_id'], 'analysis': 'The source records a payment.'}
        return httpx.Response(200, json={'choices': [{'message': {'content': json.dumps({'findings': [finding]})}}]})

    monkeypatch.setattr('app.httpx.Client', lambda **kwargs: real_client(transport=httpx.MockTransport(respond), **kwargs))
    with TestClient(create_app(settings)) as client:
        case = make_case(client); upload(client, case)
        agent = client.get('/api/cases/' + case + '/agents').json()[1]
        assert client.post('/api/cases/' + case + '/reviews', json={'agent_ids': [agent['id']], 'use_ai': True}).status_code == 409
        client.post('/api/cases/' + case + '/reviews', json={'agent_ids': [agent['id']], 'use_ai': True, 'approved_external': True})
        task = wait_for_tasks(client, case)[0]
        assert task['status'] == 'completed' and task['findings'][0]['analysis'] == 'The source records a payment.'
        invalid = True
        client.post('/api/cases/' + case + '/reviews', json={'agent_ids': [agent['id']], 'use_ai': True, 'approved_external': True})
        task = wait_for_tasks(client, case)[0]
        assert task['status'] == 'failed' and not task['findings']


def test_restart_preserves_data_and_marks_interrupted_job(tmp_path):
    settings = Settings(tmp_path)
    with TestClient(create_app(settings)) as client:
        case = make_case(client); upload(client, case)
        agent = client.get('/api/cases/' + case + '/agents').json()[0]
        with client.app.state.store.db() as db:
            db.execute('INSERT INTO tasks(id,case_id,agent_id,name,status,mode,created_at) VALUES(?,?,?,?,?,?,?)', ('interrupted', case, agent['id'], 'Interrupted', 'running', 'evidence', '2025-01-01'))
    with TestClient(create_app(settings)) as client:
        assert client.get('/api/cases').json()[0]['id'] == case
        assert len(client.get('/api/cases/' + case + '/documents').json()) == 1
        task = client.get('/api/cases/' + case + '/tasks').json()[0]
        assert task['status'] == 'failed' and 'stopped' in task['error']


def test_cross_site_writes_are_blocked(client):
    assert client.post('/api/cases', json={'name': 'Blocked'}, headers={'Sec-Fetch-Site': 'cross-site'}).status_code == 403
    assert client.post('/api/cases', json={'name': 'Blocked'}, headers={'Origin': 'https://other.example'}).status_code == 403

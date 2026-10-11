"""Sample case files made on the fly: a Word document with numbered clauses,
a table and a page break; a text PDF with running headers and footers; a
scanned PDF; a TIFF; an email with an attachment. Real Outlook .msg files are
in tests/samples/."""

import io
from email.message import EmailMessage
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from fpdf import FPDF
from PIL import Image, ImageDraw, ImageFont

SAMPLES = Path(__file__).parent / "samples"


def docx_bytes() -> bytes:
    d = Document()
    d.core_properties.title = "Witness statement of Jane Doe"
    d.add_heading("Witness statement of Jane Doe", 0)
    d.add_heading("Background", 1)
    p = d.add_paragraph("I am the Claimant. ")
    p.add_run("The agreement").bold = True
    p.add_run(" was signed on 1 March 2024.")
    d.add_paragraph("The goods were not delivered.", style="List Number")
    d.add_paragraph("I chased the Defendant twice.", style="List Number")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Date", "Event"
    t.cell(1, 0).text, t.cell(1, 1).text = "15 April 2024", "Delivery due | missed"
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    d.add_heading("Statement of truth", 1)
    d.add_paragraph("I believe that the facts stated in this witness statement are true.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


class _Claim(FPDF):
    def header(self):
        self.set_font("Helvetica", size=8)
        self.cell(0, 8, "Smith v Jones - Claim No. HC-2024-000123 - CONFIDENTIAL", align="C")
        self.ln(12)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", size=8)
        self.cell(0, 10, f"Page {self.page_no()} of 3", align="C")


def pdf_bytes() -> bytes:
    pdf = _Claim()
    for n in range(1, 4):
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 18)
        pdf.cell(0, 10, f"Part {n}: Particulars of Claim", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", size=11)
        pdf.multi_cell(0, 6, f"Paragraph {n}. The Claimant entered into a written agreement with the Defendant on 1 March 2024 for the supply of 400 tonnes of steel, to be delivered to the Claimant's site in Leeds by 15 April 2024.")
        if n == 2:
            pdf.ln(4)
            with pdf.table() as t:
                for row in (("Date", "Event", "Amount"), ("1 March 2024", "Agreement signed", "-"), ("15 April 2024", "Delivery due", "GBP 125,000")):
                    r = t.row()
                    for c in row:
                        r.cell(c)
    return bytes(pdf.output())


def _font(size: int):
    for name in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "C:/Windows/Fonts/arial.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default(size)


def scan_image(lines=("WITNESS STATEMENT", "", "My name is Jane Doe and I live", "in Leeds. The steel never arrived.")) -> Image.Image:
    img = Image.new("L", (1700, 2200), 255)
    draw = ImageDraw.Draw(img)
    font = _font(46)
    y = 220
    for line in lines:
        draw.text((160, y), line, fill=0, font=font)
        y += 80
    return img


def scanned_pdf_bytes() -> bytes:
    buf = io.BytesIO()
    scan_image().save(buf, "PNG")
    pdf = FPDF()
    pdf.add_page()
    pdf.image(io.BytesIO(buf.getvalue()), x=0, y=0, w=210)
    return bytes(pdf.output())


def tiff_bytes() -> bytes:
    buf = io.BytesIO()
    pages = [scan_image(("PAGE ONE", "", "Invoice number 4471")), scan_image(("PAGE TWO", "", "Total due 125,000"))]
    pages[0].save(buf, "TIFF", save_all=True, append_images=pages[1:])
    return buf.getvalue()


def eml_bytes(attachment: bytes | None = None) -> bytes:
    m = EmailMessage()
    m["From"] = "Jane Doe <jane@claimant.example>"
    m["To"] = "Bob Smith <bob@defendant.example>"
    m["Cc"] = "legal@claimant.example"
    m["Subject"] = "RE: Delivery of steel"
    m["Date"] = "Tue, 16 Apr 2024 09:30:00 +0100"
    m["Message-ID"] = "<abc123@claimant.example>"
    m.set_content(
        "Bob,\n\nThe steel did not arrive yesterday. Please confirm a new date.\n\nJane\n\n"
        "-----Original Message-----\nFrom: Bob Smith\nSent: 10 April 2024 16:00\nTo: Jane Doe\nSubject: Delivery of steel\n\n"
        "We are on track for delivery on 15 April.\n"
    )
    m.add_alternative("<p>Bob,</p><p>The steel did <b>not</b> arrive yesterday.</p>", subtype="html")
    if attachment is not None:
        m.add_attachment(attachment, maintype="application", subtype="pdf", filename="Delivery note.pdf")
    m.add_attachment(b"PK\x03\x04 not really a zip", maintype="application", subtype="zip", filename="photos.zip")
    return m.as_bytes()
